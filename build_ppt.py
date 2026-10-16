import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation(output_path="CampusEase_PRD_Presentation.pptx"):
    prs = Presentation()
    # Set 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    # Color Palette - Professional Slate & Modern Tech
    C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A - Primary brand dark
    C_NAVY_LIGHT = RGBColor(30, 41, 59)     # #1E293B - Card dark
    C_BLUE_PRIMARY = RGBColor(37, 99, 235)  # #2563EB - Primary action blue
    C_BLUE_LIGHT = RGBColor(239, 246, 255)  # #EFF6FF - Light blue tint
    C_SLATE_BG = RGBColor(248, 250, 252)    # #F8FAFC - Slide light bg
    C_CARD_BG = RGBColor(255, 255, 255)     # White card
    C_BORDER = RGBColor(226, 232, 240)      # #E2E8F0 - Neutral border
    C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A
    C_TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B
    C_WHITE = RGBColor(255, 255, 255)
    C_EMERALD = RGBColor(16, 185, 129)      # #10B981
    C_AMBER = RGBColor(245, 158, 11)        # #F59E0B
    C_RED = RGBColor(239, 68, 68)           # #EF4444
    C_TAG_BG = RGBColor(224, 231, 255)      # Light indigo
    C_TAG_TEXT = RGBColor(67, 56, 202)      # Deep indigo

    FONT_HEADING = "Segoe UI"
    FONT_BODY = "Segoe UI"

    def set_slide_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title, category="CAMPUSEASE PRD", is_dark=False):
        # Category Tag / Breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.name = FONT_HEADING
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_BLUE_PRIMARY if not is_dark else RGBColor(147, 197, 253)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE if is_dark else C_TEXT_MAIN

    def create_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.2)
        else:
            shape.line.fill.background()
        return shape

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s1, C_NAVY_DARK)

    # Decorative accent card behind
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.2), Inches(3.6))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = C_BLUE_PRIMARY
    accent_bar.line.fill.background()

    # Title & Subtitle box
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.8))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p = tf1.paragraphs[0]
    p.text = "CampusEase"
    p.font.name = FONT_HEADING
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "Unified Academic & Campus Operations Portal"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(147, 197, 253) # Light Blue
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.text = "Product Requirements Document (PRD) Overview • Production Release v1.0"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.space_before = Pt(20)

    # Info pills at bottom
    pills = [
        ("TARGET RELEASE", "Production v1.0", C_BLUE_PRIMARY),
        ("ARCHITECTURE", "Zero-Dependency", C_EMERALD),
        ("AUTHOR", "DeepMind Agentic Eng.", RGBColor(99, 102, 241)),
        ("STATUS", "Approved & Active", RGBColor(14, 165, 233))
    ]
    p_w = Inches(2.7)
    p_h = Inches(1.0)
    for idx, (label, val, col) in enumerate(pills):
        card = create_card(s1, Inches(0.8 + idx * 2.95), Inches(5.7), p_w, p_h, bg_color=C_NAVY_LIGHT, border_color=RGBColor(51, 65, 85))
        c_box = s1.shapes.add_textbox(Inches(0.95 + idx * 2.95), Inches(5.8), p_w - Inches(0.3), p_h - Inches(0.2))
        c_tf = c_box.text_frame
        c_tf.word_wrap = True
        c_tf.margin_top = c_tf.margin_left = c_tf.margin_right = c_tf.margin_bottom = 0
        
        pl = c_tf.paragraphs[0]
        pl.text = label
        pl.font.size = Pt(10)
        pl.font.bold = True
        pl.font.color.rgb = col

        pv = c_tf.add_paragraph()
        pv.text = val
        pv.font.size = Pt(13)
        pv.font.bold = True
        pv.font.color.rgb = C_WHITE
        pv.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: Problem Statement & Campus Challenges
    # ==========================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s2, C_SLATE_BG)
    add_header(s2, "The Problem: Fragmented Higher Education Operations", "1.1 Problem Statement")

    pain_points = [
        ("Disjointed Communication", "Notices posted on physical cork boards and unmonitored WhatsApp/Telegram chats get lost or missed, leaving students uninformed about critical deadlines.", C_RED),
        ("Manual Paper Bureaucracy", "Hostel out-passes, leave permissions, and grievances require physical slips, multiple warden signatures, and in-person footwork.", C_AMBER),
        ("Scattered Academic Resources", "Past year question papers (PYQs), timetables, and lecture notes reside in disparate personal drives or unvetted third-party repositories.", C_BLUE_PRIMARY),
        ("Zero Request Visibility", "Students lack visibility into maintenance requests (plumbing, electrical), leading to duplicate complaints and zero administrative accountability.", RGBColor(124, 58, 237))
    ]

    for idx, (title, desc, color) in enumerate(pain_points):
        x = Inches(0.8 + (idx % 2) * 5.95)
        y = Inches(1.8 + (idx // 2) * 2.5)
        card = create_card(s2, x, y, Inches(5.65), Inches(2.2), C_CARD_BG, C_BORDER)
        
        # Color indicator stripe
        stripe = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), y + Inches(0.25), Inches(0.12), Inches(0.7))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.fill.background()

        tbox = s2.shapes.add_textbox(x + Inches(0.45), y + Inches(0.2), Inches(4.9), Inches(1.8))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = C_TEXT_MAIN

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_before = Pt(8)

    # Bottom summary callout
    bot_card = create_card(s2, Inches(0.8), Inches(6.15), Inches(11.6), Inches(0.9), C_BLUE_LIGHT, RGBColor(191, 219, 254))
    bot_box = s2.shapes.add_textbox(Inches(1.1), Inches(6.25), Inches(11.0), Inches(0.7))
    bot_tf = bot_box.text_frame
    bot_tf.word_wrap = True
    bot_tf.margin_top = bot_tf.margin_left = 0
    bp = bot_tf.paragraphs[0]
    bp.text = "Result: Operational bottlenecks, administrative delays, low student satisfaction, and lost accountability."
    bp.font.bold = True
    bp.font.size = Pt(13)
    bp.font.color.rgb = C_BLUE_PRIMARY

    # ==========================================
    # SLIDE 3: Vision & Key Value Propositions
    # ==========================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s3, C_SLATE_BG)
    add_header(s3, "CampusEase: Vision & Strategic Value Propositions", "1.2 & 1.3 Vision & Values")

    # Banner Card
    ban = create_card(s3, Inches(0.8), Inches(1.7), Inches(11.7), Inches(1.2), C_NAVY_DARK, None)
    ban_box = s3.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.1), Inches(0.9))
    ban_tf = ban_box.text_frame
    ban_tf.word_wrap = True
    ban_tf.margin_top = ban_tf.margin_left = 0
    bp1 = ban_tf.paragraphs[0]
    bp1.text = "The Vision"
    bp1.font.bold = True
    bp1.font.size = Pt(12)
    bp1.font.color.rgb = RGBColor(147, 197, 253)
    bp2 = ban_tf.add_paragraph()
    bp2.text = "A lightweight, centralized institutional portal that digitizes campus operations with zero friction, complete request transparency, and unified communication."
    bp2.font.size = Pt(14)
    bp2.font.color.rgb = C_WHITE
    bp2.space_before = Pt(4)

    # 4 Pillars
    pillars = [
        ("Unified Academic Hub", "Consolidates leaves, maintenance, grievance redressal, official notices, and exam preparation resources under a single institutional login.", C_BLUE_PRIMARY),
        ("Transparent Lifecycle", "Every petition has an immutable audit log and live statuses (Pending, In Progress, Resolved, Rejected) ensuring full accountability.", C_EMERALD),
        ("Zero-Dependency Core", "Runs entirely on vanilla ES6+ frontend and Python standard library backend without heavy framework bloat or build compilation steps.", RGBColor(124, 58, 237)),
        ("Institutional Standards", "Strict accessibility compliance, distraction-free Slate Navy aesthetics, zero unnecessary animations, and strict RBAC isolation.", C_NAVY_DARK)
    ]

    for idx, (head, body, pcol) in enumerate(pillars):
        x = Inches(0.8 + idx * 2.95)
        y = Inches(3.1)
        card = create_card(s3, x, y, Inches(2.8), Inches(3.8), C_CARD_BG, C_BORDER)
        
        # Pill header
        p_strip = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.8), Inches(0.12))
        p_strip.fill.solid()
        p_strip.fill.fore_color.rgb = pcol
        p_strip.line.fill.background()

        tbox = s3.shapes.add_textbox(x + Inches(0.2), y + Inches(0.3), Inches(2.4), Inches(3.2))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_right = tf.margin_bottom = 0

        h = tf.paragraphs[0]
        h.text = f"0{idx+1}"
        h.font.size = Pt(20)
        h.font.bold = True
        h.font.color.rgb = pcol

        h2 = tf.add_paragraph()
        h2.text = head
        h2.font.size = Pt(15)
        h2.font.bold = True
        h2.font.color.rgb = C_TEXT_MAIN
        h2.space_before = Pt(6)

        b = tf.add_paragraph()
        b.text = body
        b.font.size = Pt(11)
        b.font.color.rgb = C_TEXT_MUTED
        b.space_before = Pt(10)

    # ==========================================
    # SLIDE 4: User Personas & RBAC Architecture
    # ==========================================
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s4, C_SLATE_BG)
    add_header(s4, "Role-Based Personas & Permission Boundaries", "2.0 User Personas & Roles")

    # Column 1: Student
    card_st = create_card(s4, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.3), C_CARD_BG, C_BORDER)
    st_box = s4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.1), Inches(4.8))
    st_tf = st_box.text_frame
    st_tf.word_wrap = True
    st_tf.margin_top = st_tf.margin_left = 0

    p = st_tf.paragraphs[0]
    p.text = "STUDENT PERSONA"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_BLUE_PRIMARY

    p_t = st_tf.add_paragraph()
    p_t.text = "Learners & Campus Residents"
    p_t.font.bold = True
    p_t.font.size = Pt(18)
    p_t.font.color.rgb = C_TEXT_MAIN
    p_t.space_before = Pt(2)

    p_d = st_tf.add_paragraph()
    p_d.text = "Profile: Undergrad/Postgrad hostellers & day scholars identified by Roll No, Branch & Room Number."
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = C_TEXT_MUTED
    p_d.space_before = Pt(6)

    st_actions = [
        ("Submit Digital Petitions", "Hostel out-pass, leaves, maintenance tickets, and grievances."),
        ("Real-Time Tracking", "Watch petition status transitions (Pending -> Resolved) with admin remarks."),
        ("Exam Repository Access", "Filter study materials, PYQs, and timetables by semester & branch."),
        ("Personal Task Reminders", "Set chronological academic task deadlines with automated status tags."),
        ("Campus Contacts Directory", "Access verified administrative and 24/7 emergency emergency helplines.")
    ]
    for title, desc in st_actions:
        pa = st_tf.add_paragraph()
        pa.text = f"•  {title}: "
        pa.font.bold = True
        pa.font.size = Pt(11)
        pa.font.color.rgb = C_NAVY_DARK
        pa.space_before = Pt(8)
        run = pa.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    # Column 2: Admin / Faculty
    card_ad = create_card(s4, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3), C_CARD_BG, C_BORDER)
    ad_box = s4.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    ad_tf = ad_box.text_frame
    ad_tf.word_wrap = True
    ad_tf.margin_top = ad_tf.margin_left = 0

    p = ad_tf.paragraphs[0]
    p.text = "ADMINISTRATOR & FACULTY PERSONA"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(124, 58, 237)

    p_t = ad_tf.add_paragraph()
    p_t.text = "Operations & Academic Leadership"
    p_t.font.bold = True
    p_t.font.size = Pt(18)
    p_t.font.color.rgb = C_TEXT_MAIN
    p_t.space_before = Pt(2)

    p_d = ad_tf.add_paragraph()
    p_d.text = "Profile: Wardens, Faculty Advisors, Exam Controllers & Admin Officers identified by Employee ID."
    p_d.font.size = Pt(11)
    p_d.font.color.rgb = C_TEXT_MUTED
    p_d.space_before = Pt(6)

    ad_actions = [
        ("Institutional Broadcasting", "Draft, publish, and delete notices instantly across 4 categories."),
        ("Triage & Queue Control", "Review incoming petitions, update statuses, and log review audit notes."),
        ("Backlog Maintenance", "Clear resolved tickets individually or bulk-flush completed backlogs."),
        ("Exam Cell Publishing", "Publish PYQs, syllabus, and study notes with Base64/URL attachments."),
        ("In-Place Directory Editor", "Directly update office hours, helplines, and room locations live.")
    ]
    for title, desc in ad_actions:
        pa = ad_tf.add_paragraph()
        pa.text = f"•  {title}: "
        pa.font.bold = True
        pa.font.size = Pt(11)
        pa.font.color.rgb = C_NAVY_DARK
        pa.space_before = Pt(8)
        run = pa.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 5: Core Functional Modules
    # ==========================================
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s5, C_SLATE_BG)
    add_header(s5, "Five Core Functional Pillars of CampusEase", "3.0 Functional Specifications")

    modules = [
        ("REQUESTS Service Pipeline", "Structured workflows for Leaves, Maintenance, and Academic Grievances with status transitions and immutable audit histories.", ["Hostel out-pass & leaves", "Facility defect reports", "Academic grievances"]),
        ("Notices & Bulletin Board", "Centralized chronological announcement feed categorised into Academic, Hostel, Exam, and General with real-time broadcast.", ["Category filtering", "Instant push propagation", "Admin deletion control"]),
        ("Exam Cell Repository", "Comprehensive repository for Past Year Papers (PYQs), syllabus, lecture notes, and timetables with semester/branch filters.", ["8-semester navigation", "Base64 & URL attachments", "In-browser previews"]),
        ("Personal Academic Reminders", "Isolated personal student task list with customizable deadlines, automated chronological ordering, and urgency flags.", ["'Due Soon' indicators", "Overdue status alerts", "Isolated per user ID"]),
        ("Institutional Contact Center", "Interactive directory modal displaying official emails, room locations, office hours, and emergency helplines.", ["Admin in-place editor", "Real-time updates", "Instant phone/mail links"])
    ]

    for idx, (title, summary, bullets) in enumerate(modules):
        # 3 on top row, 2 on bottom row
        if idx < 3:
            w = Inches(3.7)
            h = Inches(2.45)
            x = Inches(0.8 + idx * 3.95)
            y = Inches(1.7)
        else:
            w = Inches(5.65)
            h = Inches(2.45)
            x = Inches(0.8 + (idx - 3) * 5.95)
            y = Inches(4.45)

        card = create_card(s5, x, y, w, h, C_CARD_BG, C_BORDER)
        tbox = s5.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), w - Inches(0.5), h - Inches(0.4))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = f"{idx+1}. {title}"
        p.font.bold = True
        p.font.size = Pt(14)
        p.font.color.rgb = C_BLUE_PRIMARY

        p_s = tf.add_paragraph()
        p_s.text = summary
        p_s.font.size = Pt(10.5)
        p_s.font.color.rgb = C_TEXT_MUTED
        p_s.space_before = Pt(4)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"✔ {b}"
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = C_TEXT_MAIN
            pb.space_before = Pt(3)

    # ==========================================
    # SLIDE 6: REQUESTS Pipeline & Audit Workflow
    # ==========================================
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s6, C_SLATE_BG)
    add_header(s6, "REQUESTS Pipeline: Lifecycle, States & Audit Trail", "3.2 Service Pipeline")

    # Step indicators
    steps = [
        ("1. Submission", "Student lodges formal petition with type-specific attributes.", C_BLUE_PRIMARY),
        ("2. Pending Triage", "Queued in warden/faculty dashboard awaiting review.", C_AMBER),
        ("3. In Progress", "Under active investigation; internal notes logged.", RGBColor(124, 58, 237)),
        ("4. Final Resolution", "Approved / Resolved or Rejected with mandatory reason.", C_EMERALD)
    ]

    for idx, (st_name, st_desc, scolor) in enumerate(steps):
        x = Inches(0.8 + idx * 2.95)
        y = Inches(1.7)
        card = create_card(s6, x, y, Inches(2.8), Inches(1.4), C_CARD_BG, C_BORDER)
        
        top_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.8), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = scolor
        top_bar.line.fill.background()

        tbox = s6.shapes.add_textbox(x + Inches(0.15), y + Inches(0.18), Inches(2.5), Inches(1.1))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = 0

        p = tf.paragraphs[0]
        p.text = st_name
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = scolor

        p2 = tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_before = Pt(4)

    # Lower Two Cards: Request Types vs Admin Queue Maintenance
    # Card 1: Request Types
    c1 = create_card(s6, Inches(0.8), Inches(3.35), Inches(5.7), Inches(3.7), C_CARD_BG, C_BORDER)
    tb1 = s6.shapes.add_textbox(Inches(1.05), Inches(3.5), Inches(5.2), Inches(3.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_top = tf1.margin_left = 0

    p = tf1.paragraphs[0]
    p.text = "Petition Types & Data Schemas"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_MAIN

    types = [
        ("Hostel Leave / Out-pass", "Start date, end date, reason, emergency contact, room verification."),
        ("Facility Maintenance", "Category (Electrical, Plumbing, Furniture, Network), physical location & issue details."),
        ("Academic Grievance", "Category (Coursework, Faculty, Timetable, Exams) & formal grievance statement.")
    ]
    for t_name, t_detail in types:
        pt = tf1.add_paragraph()
        pt.text = f"• {t_name}"
        pt.font.bold = True
        pt.font.size = Pt(11)
        pt.font.color.rgb = C_BLUE_PRIMARY
        pt.space_before = Pt(8)
        pr = tf1.add_paragraph()
        pr.text = t_detail
        pr.font.size = Pt(10)
        pr.font.color.rgb = C_TEXT_MUTED
        pr.space_before = Pt(1)

    # Card 2: Queue Control & Audit History
    c2 = create_card(s6, Inches(6.8), Inches(3.35), Inches(5.7), Inches(3.7), C_CARD_BG, C_BORDER)
    tb2 = s6.shapes.add_textbox(Inches(7.05), Inches(3.5), Inches(5.2), Inches(3.3))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_top = tf2.margin_left = 0

    p = tf2.paragraphs[0]
    p.text = "Audit Logs & Queue Pruning"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_MAIN

    features = [
        ("Immutable Audit Trail", "Every status transition logs who made the update, exact timestamp, and mandatory remarks."),
        ("Single-Item Deletion", "Admin can prune erroneous or obsolete complaints on demand ('Clear This Request')."),
        ("Bulk Completed Flush", "'Clear Completed' sweeps all resolved/rejected petitions to keep operational queues lean."),
        ("Total Reset Capability", "Facilitates testing and term-end archival via authenticated 'Clear All'." )
    ]
    for f_name, f_detail in features:
        pf = tf2.add_paragraph()
        pf.text = f"✔ {f_name}: "
        pf.font.bold = True
        pf.font.size = Pt(10.5)
        pf.font.color.rgb = C_NAVY_DARK
        pf.space_before = Pt(6)
        run = pf.add_run()
        run.text = f_detail
        run.font.bold = False
        run.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 7: Exam Cell & Academic Material Repository
    # ==========================================
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s7, C_SLATE_BG)
    add_header(s7, "Exam Cell: Unified Academic Knowledge Hub", "3.4 Exam Cell Repository")

    # 4 Category Cards
    cat_items = [
        ("PYQ Repository", "Previous Year Question Papers sorted by semester, course code, and branch.", C_BLUE_PRIMARY),
        ("Curriculum & Syllabus", "Official department syllabi and marking rubrics for current semester cycles.", C_EMERALD),
        ("Lecture Notes & Handouts", "Curated notes, faculty slides, and lab instruction manuals.", RGBColor(124, 58, 237)),
        ("Exam Schedules & Timetable", "Mid-term and end-term exam date sheets with classroom seat allocations.", C_AMBER)
    ]
    for idx, (c_name, c_desc, col) in enumerate(cat_items):
        x = Inches(0.8 + idx * 2.95)
        y = Inches(1.7)
        card = create_card(s7, x, y, Inches(2.8), Inches(1.9), C_CARD_BG, C_BORDER)
        
        tbox = s7.shapes.add_textbox(x + Inches(0.18), y + Inches(0.18), Inches(2.44), Inches(1.5))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = 0

        p = tf.paragraphs[0]
        p.text = c_name
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = c_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_before = Pt(6)

    # Bottom Split Cards: Filtering Engine vs Storage & Delivery
    c_left = create_card(s7, Inches(0.8), Inches(3.85), Inches(5.7), Inches(3.2), C_CARD_BG, C_BORDER)
    tl = s7.shapes.add_textbox(Inches(1.05), Inches(4.0), Inches(5.2), Inches(2.8))
    tfl = tl.text_frame
    tfl.word_wrap = True
    tfl.margin_top = tfl.margin_left = 0

    p = tfl.paragraphs[0]
    p.text = "Fast Departmental Filtering Engine"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_MAIN

    f_points = [
        ("8-Semester Granular Navigation", "Immediate instant-filtering across Semesters 1 through 8."),
        ("Multi-Branch Segregation", "CSE, ECE, Mechanical, Civil, Chemical, Electrical, and Sciences."),
        ("Zero Latency In-Memory Query", "Frontend instantly filters without full page reloads.")
    ]
    for h, b in f_points:
        p = tfl.add_paragraph()
        p.text = f"• {h}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = b
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED

    c_right = create_card(s7, Inches(6.8), Inches(3.85), Inches(5.7), Inches(3.2), C_CARD_BG, C_BORDER)
    tr = s7.shapes.add_textbox(Inches(7.05), Inches(4.0), Inches(5.2), Inches(2.8))
    tfr = tr.text_frame
    tfr.word_wrap = True
    tfr.margin_top = tfr.margin_left = 0

    p = tfr.paragraphs[0]
    p.text = "Attachment & Offline Delivery Pipeline"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_TEXT_MAIN

    d_points = [
        ("Base64 Local Storage Bridge", "Facilitates file uploads directly embedded into portable JSON storage."),
        ("Remote URL Hyperlinking", "Support for cloud storage, Google Drive links, and CDN assets."),
        ("One-Click Offline Download", "Students can download question papers directly for offline exam prep."),
        ("In-Browser Image Previews", "Instant previewing of diagrams, tables, and notices.")
    ]
    for h, b in d_points:
        p = tfr.add_paragraph()
        p.text = f"✔ {h}: "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_EMERALD
        p.space_before = Pt(6)
        r = p.add_run()
        r.text = b
        r.font.bold = False
        r.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 8: System Architecture & Zero-Dependency Stack
    # ==========================================
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s8, C_SLATE_BG)
    add_header(s8, "System Architecture: Zero-Dependency Philosophy", "4.0 Technical Stack")

    tech_layers = [
        ("Client Layer (SPA)", "HTML5 + Modern CSS + Vanilla ES6+", [
            "Instant page load (<200ms)",
            "Zero node_modules / webpack bloat",
            "PWA web manifest enabled",
            "100% portable single-page app"
        ], C_BLUE_PRIMARY),
        ("Backend Layer (Python Core)", "Python 3.10+ Standard Library", [
            "Built with http.server & urllib",
            "Zero pip install requirements",
            "Dual WSGI & BaseHTTPRequestHandler",
            "Native JSON serialization"
        ], C_EMERALD),
        ("Storage Layer (Document DB)", "Atomic JSON Document Store", [
            "Human-readable database.json",
            "/tmp runtime bridge for serverless",
            "Atomic read/write operations",
            "Zero external DB provisioning"
        ], RGBColor(124, 58, 237)),
        ("Deployment Target", "Multi-Cloud & Serverless Ready", [
            "Vercel Serverless (vercel.json)",
            "Render / Railway (Procfile)",
            "Local standalone script",
            "Zero vendor lock-in"
        ], C_NAVY_DARK)
    ]

    for idx, (layer_name, sub, bullets, lcol) in enumerate(tech_layers):
        x = Inches(0.8 + idx * 2.95)
        y = Inches(1.7)
        card = create_card(s8, x, y, Inches(2.8), Inches(4.5), C_CARD_BG, C_BORDER)
        
        top_bar = s8.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(2.8), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = lcol
        top_bar.line.fill.background()

        tbox = s8.shapes.add_textbox(x + Inches(0.18), y + Inches(0.25), Inches(2.44), Inches(4.1))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = 0

        p = tf.paragraphs[0]
        p.text = layer_name
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = C_TEXT_MAIN

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.bold = True
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = lcol
        p_sub.space_before = Pt(3)

        for b in bullets:
            pb = tf.add_paragraph()
            pb.text = f"• {b}"
            pb.font.size = Pt(10)
            pb.font.color.rgb = C_TEXT_MUTED
            pb.space_before = Pt(8)

    # Footnote bar
    foot = create_card(s8, Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.7), C_BLUE_LIGHT, RGBColor(191, 219, 254))
    ft_box = s8.shapes.add_textbox(Inches(1.0), Inches(6.42), Inches(11.2), Inches(0.55))
    ft_tf = ft_box.text_frame
    ft_tf.word_wrap = True
    ft_tf.margin_top = ft_tf.margin_left = 0
    fp = ft_tf.paragraphs[0]
    fp.text = "Key Advantage: Anyone with Python installed can run the entire system instantly using a single command: python server.py"
    fp.font.bold = True
    fp.font.size = Pt(11.5)
    fp.font.color.rgb = C_BLUE_PRIMARY

    # ==========================================
    # SLIDE 9: REST API Interface Specifications
    # ==========================================
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s9, C_SLATE_BG)
    add_header(s9, "REST API Endpoints & Route Dispatcher", "5.0 API Interface Specification")

    endpoints = [
        ("POST /api/login & /api/register", "Authentication", "RBAC session initiation for students (Roll No) and faculty/admin (Emp ID).", C_BLUE_PRIMARY),
        ("GET /api/notices", "Notices", "Fetch institutional notices with category filtering (Academic, Hostel, Exam, General).", C_NAVY_DARK),
        ("POST & DELETE /api/notices/{id}", "Notices (Admin)", "Publish announcements or delete obsolete circulars with instant broadcast.", C_RED),
        ("GET /api/requests?who={name}&role={role}", "Requests", "Filter petitions by student ownership or view complete admin triage queue.", C_BLUE_PRIMARY),
        ("POST /api/requests", "Requests", "Submit new Leave Application, Facility Maintenance ticket, or Academic Grievance.", C_EMERALD),
        ("PATCH & DELETE /api/requests/{id}", "Requests", "Update status (Pending -> Resolved) + append history note; delete single record.", RGBColor(124, 58, 237)),
        ("GET & POST /api/exam", "Exam Cell", "Query study resources by semester/branch; upload question papers with media.", C_AMBER),
        ("GET & POST /api/contact", "Directory", "Fetch contact metadata; update phone, email, office hours, and helplines in-place.", C_NAVY_DARK)
    ]

    for idx, (route, cat, desc, col) in enumerate(endpoints):
        col_idx = idx % 2
        row_idx = idx // 2
        x = Inches(0.8 + col_idx * 5.95)
        y = Inches(1.7 + row_idx * 1.3)
        w = Inches(5.65)
        h = Inches(1.15)

        card = create_card(s9, x, y, w, h, C_CARD_BG, C_BORDER)
        
        # Method color indicator
        ind = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
        ind.fill.solid()
        ind.fill.fore_color.rgb = col
        ind.line.fill.background()

        tbox = s9.shapes.add_textbox(x + Inches(0.2), y + Inches(0.12), w - Inches(0.35), h - Inches(0.2))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = 0

        p = tf.paragraphs[0]
        p.text = route
        p.font.name = "Consolas"
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_NAVY_DARK

        p_desc = tf.add_paragraph()
        p_desc.text = f"[{cat}] {desc}"
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = C_TEXT_MUTED
        p_desc.space_before = Pt(3)

    # ==========================================
    # SLIDE 10: Non-Functional Requirements & Aesthetics
    # ==========================================
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s10, C_SLATE_BG)
    add_header(s10, "Non-Functional Requirements & Design Principles", "6.0 NFR & Design System")

    # 3 Large Metric Cards
    metrics = [
        ("< 0.8s", "First Contentful Paint (FCP)", "Ultra-fast load time even over constrained 4G campus Wi-Fi networks.", C_BLUE_PRIMARY),
        ("< 50ms", "API Turnaround Time", "Sub-50 millisecond response on local server or edge serverless runtimes.", C_EMERALD),
        ("< 100 KB", "Total Static Bundle Size", "Featherweight codebase excluding user media payloads.", RGBColor(124, 58, 237))
    ]
    for idx, (num, label, desc, col) in enumerate(metrics):
        x = Inches(0.8 + idx * 3.95)
        y = Inches(1.7)
        card = create_card(s10, x, y, Inches(3.7), Inches(2.0), C_CARD_BG, C_BORDER)
        
        tbox = s10.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), Inches(3.2), Inches(1.6))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_left = 0

        p = tf.paragraphs[0]
        p.text = num
        p.font.bold = True
        p.font.size = Pt(28)
        p.font.color.rgb = col

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.bold = True
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_MAIN
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = C_TEXT_MUTED
        p3.space_before = Pt(4)

    # 2 Comparison Boxes: Institutional Aesthetics vs Anti-Patterns
    c_left = create_card(s10, Inches(0.8), Inches(3.95), Inches(5.7), Inches(3.1), C_CARD_BG, C_BORDER)
    tl = s10.shapes.add_textbox(Inches(1.05), Inches(4.1), Inches(5.2), Inches(2.8))
    tfl = tl.text_frame
    tfl.word_wrap = True
    tfl.margin_top = tfl.margin_left = 0

    p = tfl.paragraphs[0]
    p.text = "Institutional Design Standards"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_EMERALD

    dos = [
        "Slate Navy (#0F172A) & crisp neutral borders (#E2E8F0)",
        "Strict 6px border-radius standard on all buttons & cards",
        "Semantic SVG line icons for universal readability",
        "Subtle semantic status badges (Green, Amber, Red)"
    ]
    for d in dos:
        p = tfl.add_paragraph()
        p.text = f"✔ {d}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_MAIN
        p.space_before = Pt(6)

    c_right = create_card(s10, Inches(6.8), Inches(3.95), Inches(5.7), Inches(3.1), C_CARD_BG, C_BORDER)
    tr = s10.shapes.add_textbox(Inches(7.05), Inches(4.1), Inches(5.2), Inches(2.8))
    tfr = tr.text_frame
    tfr.word_wrap = True
    tfr.margin_top = tfr.margin_left = 0

    p = tfr.paragraphs[0]
    p.text = "Strict Anti-Patterns (Eliminated)"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_RED

    donts = [
        "No distracting purple neon gradients or playful hover flips",
        "No pill-shaped buttons or whimsical floating elements",
        "No decorative emojis; strictly professional UI symbols",
        "No pre-filled or mock form credentials; zero fake reviews"
    ]
    for d in donts:
        p = tfr.add_paragraph()
        p.text = f"✖ {d}"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_MUTED
        p.space_before = Pt(6)

    # ==========================================
    # SLIDE 11: Implementation Status & Roadmap
    # ==========================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s11, C_SLATE_BG)
    add_header(s11, "Implementation Status & Future Roadmap", "7.0 Product Roadmap")

    # Column 1: Completed in v1.0
    c_v1 = create_card(s11, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.3), C_CARD_BG, C_BORDER)
    tb_v1 = s11.shapes.add_textbox(Inches(1.05), Inches(1.9), Inches(5.2), Inches(4.9))
    tf_v1 = tb_v1.text_frame
    tf_v1.word_wrap = True
    tf_v1.margin_top = tf_v1.margin_left = 0

    p = tf_v1.paragraphs[0]
    p.text = "VERSION 1.0 (PRODUCTION COMPLETED)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_EMERALD

    p_t = tf_v1.add_paragraph()
    p_t.text = "Core Portal Capabilities"
    p_t.font.bold = True
    p_t.font.size = Pt(18)
    p_t.font.color.rgb = C_TEXT_MAIN
    p_t.space_before = Pt(2)

    v1_items = [
        "Student & Admin Authentication & RBAC Isolation",
        "Notice Board with Category Filters & Instant Broadcast",
        "REQUESTS Pipeline with Full Audit History & Notes",
        "Admin Queue Management (Single & Bulk Ticket Clearing)",
        "Personal Student Reminders with Due Date Sorting",
        "Exam Cell Repository with Media Upload & Download",
        "Dynamic Admin-Editable Institutional Contact Modal",
        "Vercel Serverless Entrypoint & Edge Read/Write Support"
    ]
    for item in v1_items:
        pi = tf_v1.add_paragraph()
        pi.text = f"✔  {item}"
        pi.font.size = Pt(10.5)
        pi.font.color.rgb = C_TEXT_MAIN
        pi.space_before = Pt(5)

    # Column 2: Roadmap for v2.0 & v2.1
    c_v2 = create_card(s11, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.3), C_CARD_BG, C_BORDER)
    tb_v2 = s11.shapes.add_textbox(Inches(7.05), Inches(1.9), Inches(5.2), Inches(4.9))
    tf_v2 = tb_v2.text_frame
    tf_v2.word_wrap = True
    tf_v2.margin_top = tf_v2.margin_left = 0

    p = tf_v2.paragraphs[0]
    p.text = "VERSIONS 2.0 & 2.1 (FUTURE ROADMAP)"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = C_BLUE_PRIMARY

    p_t = tf_v2.add_paragraph()
    p_t.text = "Enterprise Expansion"
    p_t.font.bold = True
    p_t.font.size = Pt(18)
    p_t.font.color.rgb = C_TEXT_MAIN
    p_t.space_before = Pt(2)

    v2_items = [
        ("Persistent Relational Database Migration", "Transition from JSON store to PostgreSQL / SQLite with migration utilities for high concurrency.", "v2.0 Planned"),
        ("Automated Email & SMS Notifications", "Instant automated dispatch to parents/wardens upon leave grant or ticket resolution.", "v2.0 Planned"),
        ("Multi-Campus Role Hierarchies", "Multi-tiered administration across Deans, HODs, Wardens, and Lab Assistants.", "v2.1 Planned"),
        ("Mobile PWA Offline Sync", "Background sync for offline question paper browsing and ticket drafting.", "v2.1 Planned")
    ]
    for title, desc, tag in v2_items:
        pi = tf_v2.add_paragraph()
        pi.text = f"➤  {title} "
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = C_NAVY_DARK
        pi.space_before = Pt(8)
        
        pr = tf_v2.add_paragraph()
        pr.text = f"    {desc}"
        pr.font.size = Pt(9.5)
        pr.font.color.rgb = C_TEXT_MUTED
        pr.space_before = Pt(1)

    # ==========================================
    # SLIDE 12: Executive Summary & Conclusion (Dark Theme)
    # ==========================================
    s12 = prs.slides.add_slide(blank_slide_layout)
    set_slide_bg(s12, C_NAVY_DARK)

    # Center card / box
    cnt_box = s12.shapes.add_textbox(Inches(1.5), Inches(1.0), Inches(10.3), Inches(1.6))
    tf12 = cnt_box.text_frame
    tf12.word_wrap = True
    tf12.margin_top = tf12.margin_left = 0

    p = tf12.paragraphs[0]
    p.text = "CAMPUSEASE"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(147, 197, 253)

    p2 = tf12.add_paragraph()
    p2.text = "Transforming Campus Operations into an Effortless Experience"
    p2.font.bold = True
    p2.font.size = Pt(26)
    p2.font.color.rgb = C_WHITE
    p2.space_before = Pt(4)

    # 3 Summary Pillar Cards
    summary_cards = [
        ("Zero Friction", "Instant setup with no dependencies, lightweight footprint, and multi-cloud readiness.", C_BLUE_PRIMARY),
        ("Total Accountability", "Complete transparency into leaves, grievances, and maintenance with timestamped audit trails.", C_EMERALD),
        ("Institutional Grade", "Professional, distraction-free aesthetic built to strict accessibility and RBAC standards.", RGBColor(124, 58, 237))
    ]

    for idx, (title, desc, col) in enumerate(summary_cards):
        x = Inches(1.5 + idx * 3.55)
        y = Inches(3.0)
        card = create_card(s12, x, y, Inches(3.2), Inches(2.3), C_NAVY_LIGHT, RGBColor(51, 65, 85))
        
        cbox = s12.shapes.add_textbox(x + Inches(0.25), y + Inches(0.25), Inches(2.7), Inches(1.8))
        ctf = cbox.text_frame
        ctf.word_wrap = True
        ctf.margin_top = ctf.margin_left = 0

        cp = ctf.paragraphs[0]
        cp.text = title
        cp.font.bold = True
        cp.font.size = Pt(16)
        cp.font.color.rgb = col

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = RGBColor(203, 213, 225)
        cp2.space_before = Pt(8)

    # Bottom Contact & Q&A text
    bot_box = s12.shapes.add_textbox(Inches(1.5), Inches(5.8), Inches(10.3), Inches(0.9))
    b_tf = bot_box.text_frame
    b_tf.word_wrap = True
    b_tf.margin_top = b_tf.margin_left = 0

    bp = b_tf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = "Production v1.0 Ready for Deployment"
    bp.font.bold = True
    bp.font.size = Pt(15)
    bp.font.color.rgb = C_WHITE

    bp2 = b_tf.add_paragraph()
    bp2.alignment = PP_ALIGN.CENTER
    bp2.text = "CampusEase: Empowering students, faculty, and administration with a single source of truth."
    bp2.font.size = Pt(12)
    bp2.font.color.rgb = RGBColor(148, 163, 184)
    bp2.space_before = Pt(4)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation generated successfully at: {output_path}")

if __name__ == "__main__":
    build_presentation()
